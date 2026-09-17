from pathlib import Path
from copy import copy
from openpyxl import Workbook

root = Path(__file__).resolve().parent.parent
out = root / ".work" / "datacamp-audit-2026-09-17" / "ai-finance-synthetic-resources"
out.mkdir(parents=True, exist_ok=True)

(out / "harborview_earnings_excerpt_neopolis_synthetique.txt").write_text("""DONNÉES SYNTHÉTIQUES NEOPOLIS — EXERCICE UNIQUEMENT

Orion Technologies — extrait fictif de conférence téléphonique

Direction : l’équipe a simplifié la coordination entre produits et priorités de livraison. Elle vise une communication plus régulière sur les jalons à venir.
Stratégie : les responsables privilégient une allocation prudente des ressources et un suivi mensuel des risques opérationnels.
Risques : les hypothèses de planification restent soumises à l’évolution de la demande et à la disponibilité des équipes.

Cet extrait est fictif, non financier et ne constitue ni une recommandation d’investissement ni un document d’entreprise.
""", encoding="utf-8")

(out / "finwise_brand_guidelines_neopolis_synthetique.txt").write_text("""LIGNES DIRECTRICES SYNTHÉTIQUES NEOPOLIS — EXERCICE UNIQUEMENT

Marque : FinWise, scénario pédagogique fictif.
Rôle attendu : aider à structurer une synthèse de recherche à partir d’un contenu pédagogique.
Ton : professionnel, nuancé, clair et sans sensationnalisme.
Format : commencer par le contexte, distinguer faits et hypothèses, terminer par des limites et questions à vérifier.
Sécurité : ne pas utiliser de données clients, de données de marché non publiées, de données personnelles ou de recommandations d’investissement.

Ce document est fictif et ne représente aucune entreprise, directive ou stratégie réelle.
""", encoding="utf-8")

wb = Workbook()
ws = wb.active
ws.title = "Synthetic Forecast"
ws.append(["Period", "Revenue", "Operating expense", "Scenario note"])
ws.append(["Q1", 100, 62, "Baseline synthetic scenario"])
ws.append(["Q2", 108, 66, "Baseline synthetic scenario"])
ws.append(["Q3", 112, 70, "Baseline synthetic scenario"])
ws.append(["Q4", 118, 73, "Illustrative planning scenario"])
for cell in ws[1]:
    font = copy(cell.font)
    font.bold = True
    cell.font = font
ws.column_dimensions["A"].width = 15
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 22
ws.column_dimensions["D"].width = 34
notes = wb.create_sheet("Read me")
notes["A1"] = "SYNTHETIC DATA — LEARNING EXERCISE ONLY"
notes["A2"] = "All figures are fictional illustrative units. Do not use for financial decisions."
notes.column_dimensions["A"].width = 92
wb.save(out / "q3_forecast_data_neopolis_synthetique.xlsx")
print(out)
